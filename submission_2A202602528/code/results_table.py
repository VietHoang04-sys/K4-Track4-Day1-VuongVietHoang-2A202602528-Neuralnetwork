"""Persist experiment results and populate the provided workbook template.

Nhiệm vụ: lưu kết quả từng lần chạy ra JSON, rồi điền vào experiments.xlsx từ mẫu
templates/experiment_table_template.xlsx (đừng gõ tay hàng chục dòng, rất dễ sai).

Tên cột của sheet "Experiments" (giữ nguyên, đúng thứ tự mẫu):
    exp_id, group, description, loss, optimizer, lr, weight_decay, batch, epochs, hidden, dropout,
    clip_norm, precision, init, seed, step0_loss, best_val_loss, best_epoch, final_train_loss,
    final_val_loss, val_acc, val_macro_f1, time_per_epoch_s, peak_mem_MB, diverged,
    eval_acc, eval_macro_f1, figure_file, notes
(các cột công thức ở cuối bảng mẫu tự tính, đừng ghi đè)
"""
from __future__ import annotations

import json
import math
from pathlib import Path

FORMULA_COLUMNS = {
    "step0_gap_vs_lnC",
    "gap_val_minus_train",
    "delta_val_f1_vs_base",
    "beyond_noise",
}
EXPERIMENT_COLUMNS = [
    "exp_id", "group", "description", "loss", "optimizer", "lr",
    "weight_decay", "batch", "epochs", "hidden", "dropout", "clip_norm",
    "precision", "init", "seed", "step0_loss", "best_val_loss",
    "best_epoch", "final_train_loss", "final_val_loss", "val_acc",
    "val_macro_f1", "time_per_epoch_s", "peak_mem_MB", "diverged",
    "eval_acc", "eval_macro_f1", "figure_file", "notes",
]

def save_result(result: dict, results_dir: str = "../results") -> str:
    """Ghi result["cfg"], result["history"], result["summary"] (KHÔNG ghi best_state) ra
    <results_dir>/<exp_id>.json. Trả về đường dẫn file. Tạo thư mục nếu chưa có."""
    exp_id = result["cfg"]["exp_id"]
    output = Path(results_dir)
    output.mkdir(parents=True, exist_ok=True)
    payload = {key: value for key, value in result.items() if key != "best_state"}
    path = output / f"{exp_id}.json"
    with path.open("w", encoding="utf-8") as stream:
        json.dump(payload, stream, ensure_ascii=False, indent=2, allow_nan=False)
    return str(path)


def load_results(results_dir: str = "../results") -> list[dict]:
    """Đọc mọi file *.json trong results_dir, trả về danh sách dict (sắp theo exp_id)."""
    directory = Path(results_dir)
    if not directory.is_dir():
        return []
    results = []
    for path in sorted(directory.glob("*.json")):
        with path.open(encoding="utf-8") as stream:
            results.append(json.load(stream))
    return results


def to_row(result: dict, eval_scores: dict | None = None, notes: str = "") -> dict:
    """Biến một kết quả thành một dòng của bảng: gộp cfg + summary (+ eval_acc, eval_macro_f1 nếu có)
    + figure_file = f"figures/{exp_id}.png". Khoá phải trùng tên cột ở đầu file.
    Chỉ truyền eval_scores cho baseline và cấu hình cuối cùng."""
    cfg = result["cfg"]
    summary = result["summary"]
    row = {column: None for column in EXPERIMENT_COLUMNS}
    for column in EXPERIMENT_COLUMNS:
        if column in cfg:
            row[column] = cfg[column]
        if column in summary:
            row[column] = summary[column]
    row["hidden"] = "-".join(str(width) for width in cfg["hidden"])
    row["figure_file"] = f"figures/{cfg['exp_id']}.png"
    row["notes"] = notes
    if eval_scores is not None:
        row["eval_acc"] = eval_scores["accuracy"]
        row["eval_macro_f1"] = eval_scores["macro_f1"]
    return row


def write_xlsx(rows: list[dict], template_path: str, out_path: str) -> None:
    """Điền các dòng vào sheet "Experiments" của mẫu, từ dòng 2 trở xuống, rồi lưu thành out_path.

    Các bước (openpyxl):
      1. wb = openpyxl.load_workbook(template_path)   # KHÔNG dùng data_only=True (sẽ mất công thức)
      2. ws = wb["Experiments"]; đọc tiêu đề dòng 1 để biết cột nào ứng với khoá nào
      3. với mỗi row: ghi giá trị vào đúng cột; BỎ QUA các cột công thức (step0_gap_vs_lnC, gap_val_minus_train,
         delta_val_f1_vs_base, beyond_noise)
      4. wb.save(out_path)
    Sau khi lưu, mở file bằng Excel/LibreOffice để các công thức tính lại.
    """
    from openpyxl import load_workbook

    template = Path(template_path)
    if not template.is_file():
        raise FileNotFoundError(f"Workbook template not found: {template}")
    if len(rows) != len({row.get("exp_id") for row in rows}):
        raise ValueError("Each experiment row must have a unique exp_id.")
    workbook = load_workbook(template)
    if "Experiments" not in workbook.sheetnames:
        raise KeyError("Workbook template must contain an 'Experiments' sheet.")
    sheet = workbook["Experiments"]
    headers = [sheet.cell(1, column).value for column in range(1, sheet.max_column + 1)]
    unknown = set().union(*(set(row) for row in rows)) - set(headers)
    if unknown:
        raise ValueError(f"Unexpected experiment columns: {sorted(unknown)}")

    first_data_row = 2
    required_last_row = first_data_row + len(rows) - 1
    if required_last_row > sheet.max_row:
        start = sheet.max_row + 1
        sheet.insert_rows(start, required_last_row - sheet.max_row)
        for row_index in range(start, required_last_row + 1):
            for column, header in enumerate(headers, start=1):
                if header in FORMULA_COLUMNS:
                    continue
                template_cell = sheet.cell(2, column)
                target_cell = sheet.cell(row_index, column)
                if template_cell.has_style:
                    target_cell._style = template_cell._style
                if template_cell.number_format:
                    target_cell.number_format = template_cell.number_format

    for row_index in range(first_data_row, sheet.max_row + 1):
        for column, header in enumerate(headers, start=1):
            cell = sheet.cell(row_index, column)
            if header in FORMULA_COLUMNS:
                if row_index >= first_data_row + len(rows):
                    cell.value = None
                elif row_index > 2:
                    template_formula = sheet.cell(2, column).value
                    if isinstance(template_formula, str) and template_formula.startswith("="):
                        from openpyxl.formula.translate import Translator

                        cell.value = Translator(
                            template_formula, origin=sheet.cell(2, column).coordinate
                        ).translate_formula(cell.coordinate)
                continue
            cell.value = None

    for offset, row in enumerate(rows):
        row_index = first_data_row + offset
        for column, header in enumerate(headers, start=1):
            if header in FORMULA_COLUMNS:
                continue
            value = row.get(header)
            if isinstance(value, float) and not math.isfinite(value):
                value = None
            sheet.cell(row_index, column, value)

    output = Path(out_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output)
